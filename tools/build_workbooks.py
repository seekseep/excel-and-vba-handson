#!/usr/bin/env python3
"""
各レクチャーの book/ に置く練習用ブックを生成する。

  - .xlsx … マクロなし。節を始める状態。本文の表計算デモが描き、章の最初の節ではそのまま配る。
           同じ章で中身が変わらない節には置かない（デモは src= で最初の節のブックを指す）
  - .xlsm … .xlsx と同じ中身に、その節を終えた状態のマクロ（example/vba/*.bas）を入れたもの。
           Excel の受講者がダウンロードするのはこちら（build-downloads.mjs）
スプレッドシートの受講者は、Google ドライブのサンプルをコピーして始める（sections/spreadsheets.json）。
.xlsm のマクロの部分（vbaProject.bin）は tools/vba_project.py が仕様どおりに組み立てる（Excel は不要）。

使い方（保守者のみ。受講者は実行しない）:
    pip install openpyxl
    python3 tools/build_workbooks.py

売上データは乱数の種を固定しているので、何度実行しても同じ中身になる。
"""

import random
import zipfile
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from vba_project import build_vba_project

ROOT = Path(__file__).resolve().parent.parent
SECTIONS = ROOT / "sections"

SHOP = {
    "name": "そよかぜ文具店",
    "zip": "〒150-0000",
    "address": "東京都渋谷区そよかぜ町 1-2-3",
    "tel": "TEL 03-0000-0000",
}

CUSTOMERS = [
    # 顧客名, 郵便番号, 住所, ご担当者
    ("ひまわり商事", "〒104-0000", "東京都中央区ひまわり 2-4-6", "総務部 山田 様"),
    ("あおば設計事務所", "〒160-0000", "東京都新宿区あおば台 3-1-8", "事務担当 佐藤 様"),
    ("みなと物産", "〒231-0000", "神奈川県横浜市中区みなと 5-5-1", "経理部 鈴木 様"),
    ("さくら学習塾", "〒180-0000", "東京都武蔵野市さくら通り 1-9", "教室長 高橋 様"),
    ("こもれび珈琲", "〒154-0000", "東京都世田谷区こもれび 7-2-10", "店長 田中 様"),
]

PRODUCTS = [
    # 商品, 単価
    ("ボールペン", 120),
    ("ノート", 180),
    ("クリアファイル", 60),
    ("付箋", 250),
    ("コピー用紙", 550),
    ("ホッチキス", 680),
    ("修正テープ", 320),
]

SALES_ROWS = 30

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '#,##0'
DATE = 'yyyy/mm/dd'


def make_sales():
    """2026 年 10 月の売上明細 30 行。日付順に並べる。"""
    rng = random.Random(20261001)
    rows = []
    for _ in range(SALES_ROWS):
        day = rng.randint(1, 30)
        customer = rng.choice(CUSTOMERS)[0]
        product, price = rng.choice(PRODUCTS)
        quantity = rng.choice([1, 2, 3, 5, 10, 12, 20])
        rows.append((date(2026, 10, day), customer, product, price, quantity))
    rows.sort(key=lambda r: (r[0], r[1]))
    return rows


def write_sales_sheet(ws, sales, with_amount):
    headers = ["日付", "顧客", "商品", "単価", "数量", "金額"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOX
    for r, (d, customer, product, price, quantity) in enumerate(sales, start=2):
        values = [d, customer, product, price, quantity, price * quantity if with_amount else None]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = BOX
        ws.cell(row=r, column=1).number_format = DATE
        ws.cell(row=r, column=4).number_format = YEN
        ws.cell(row=r, column=6).number_format = YEN
    widths = {"A": 12, "B": 18, "C": 16, "D": 9, "E": 7, "F": 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def write_summary_area(ws):
    """3 章の集計欄。I 列に関数、J 列に VBA の結果を並べて比べる。"""
    def header(cell_ref, value):
        cell = ws[cell_ref]
        cell.value = value
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOX

    header("H1", "項目")
    header("I1", "関数")
    header("J1", "VBA")
    ws["H2"] = "売上合計"
    ws["H2"].border = BOX
    for ref in ("I2", "J2"):
        ws[ref].border = BOX
        ws[ref].number_format = YEN
    ws["I2"].fill = INPUT_FILL

    header("H4", "顧客")
    header("I4", "関数")
    header("J4", "VBA")
    for i, (name, *_rest) in enumerate(CUSTOMERS, start=5):
        ws.cell(row=i, column=8, value=name).border = BOX
        for col in (9, 10):
            cell = ws.cell(row=i, column=col)
            cell.border = BOX
            cell.number_format = YEN
        ws.cell(row=i, column=9).fill = INPUT_FILL

    ws.column_dimensions["G"].width = 3
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 12


def write_customer_sheet(ws):
    headers = ["顧客名", "郵便番号", "住所", "ご担当者"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOX
    for r, row in enumerate(CUSTOMERS, start=2):
        for col, v in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=v).border = BOX
    for col, w in {"A": 18, "B": 12, "C": 32, "D": 18}.items():
        ws.column_dimensions[col].width = w


def write_invoice_template(ws):
    """
    請求書のひな形。マクロはこのシートをコピーして値を埋める。
    セル番地は VBA / GAS のコードと対応しているので、動かすときは両方を直すこと。
      A5 宛名 / A6 住所 / A7 ご担当者 / E3 発行日 / B9 ご請求金額
      12〜31 行目 明細（A 日付, B 商品, C 単価, D 数量, E 金額） / E33 合計
    """
    for col, w in {"A": 14, "B": 22, "C": 12, "D": 10, "E": 14}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:E1")
    ws["A1"] = "請 求 書"
    ws["A1"].font = Font(size=20, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["D3"] = "発行日"
    ws["D3"].alignment = Alignment(horizontal="right")
    ws["E3"].number_format = DATE

    ws["A5"] = "（宛名）"
    ws["A5"].font = Font(size=14, bold=True)
    ws["A6"] = "（住所）"
    ws["A7"] = "（ご担当者）"

    ws["D5"] = SHOP["name"]
    ws["D5"].font = BOLD
    ws["D6"] = SHOP["zip"]
    ws["D7"] = SHOP["address"]
    ws["D8"] = SHOP["tel"]

    ws["A9"] = "ご請求金額"
    ws["A9"].font = BOLD
    ws["B9"].number_format = '"¥"#,##0'
    ws["B9"].font = Font(size=14, bold=True)
    ws["B9"].border = Border(bottom=Side(style="medium"))

    headers = ["日付", "商品", "単価", "数量", "金額"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
    for r in range(12, 32):
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BOX
        ws.cell(row=r, column=1).number_format = DATE
        ws.cell(row=r, column=3).number_format = YEN
        ws.cell(row=r, column=5).number_format = YEN

    ws["D33"] = "合計"
    ws["D33"].font = BOLD
    ws["D33"].border = BOX
    ws["E33"].number_format = YEN
    ws["E33"].font = BOLD
    ws["E33"].border = BOX

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:E33"


def save(wb, lecture_dir, name, end_state=None, buttons=(), start_buttons=None, with_start=True):
    """
    .xlsx（節を始める状態。本文のデモもこれを描く）を保存し、マクロ入りの .xlsm も作る。
    with_start=False なら .xlsx は作らない（同じ章の前の節と同じ中身のとき）。
    end_state(wb) を渡すと、.xlsm はその節のマクロを実行し終えた後の値まで入れた状態にする
    （次の節を、ダウンロードした .xlsm からそのまま始められるように）。
    buttons はマクロを実行するボタン（Button の並び）。start_buttons を渡すと .xlsx にはそちらを置く
    （5-3 のように、ボタンを自分で置く節のため）。
    """
    if with_start:
        path = lecture_dir / "book" / f"{name}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        add_buttons(wb, path, buttons if start_buttons is None else start_buttons)
        print(f"[build_workbooks] {path.relative_to(ROOT)}")
    if end_state:
        end_state(wb)
    save_xlsm(wb, lecture_dir, name, buttons)


# ---------------------------------------------------------------------------
# マクロを実行するボタン（フォーム コントロールのボタン）
# ---------------------------------------------------------------------------

class Button:
    """sheet の cell を左上に、cols 列 × rows 行の大きさで置く。押すと macro が動く。"""

    def __init__(self, sheet, cell, text, macro, cols=2, rows=2):
        self.sheet = sheet
        self.cell = cell
        self.text = text
        self.macro = macro
        self.cols = cols
        self.rows = rows


VML_HEAD = (
    '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
    ' xmlns:x="urn:schemas-microsoft-com:office:excel">'
    '<o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="{idmap}"/></o:shapelayout>'
    '<v:shapetype id="_x0000_t201" coordsize="21600,21600" o:spt="201" path="m,l,21600r21600,l21600,xe">'
    '<v:stroke joinstyle="miter"/>'
    '<v:path shadowok="f" o:extrusionok="f" strokeok="f" fillok="f" o:connecttype="rect"/>'
    '<o:lock v:ext="edit" shapetype="t"/>'
    '</v:shapetype>'
)
VML_BUTTON = (
    '<v:shape id="_x0000_s{shape_id}" type="#_x0000_t201"'
    ' style="position:absolute;margin-left:0;margin-top:0;width:96pt;height:28pt;z-index:{z};mso-wrap-style:tight"'
    ' o:button="t" fillcolor="buttonFace [67]" strokecolor="windowText [64]" o:insetmode="auto">'
    '<v:fill color2="buttonFace [67]" o:detectmouseclick="t"/>'
    '<o:lock v:ext="edit" rotation="t"/>'
    '<v:textbox style="mso-direction-alt:auto" o:singleclick="f">'
    '<div style="text-align:center"><font face="游ゴシック" size="220" color="#000000">{text}</font></div>'
    '</v:textbox>'
    '<x:ClientData ObjectType="Button">'
    '<x:Anchor>{anchor}</x:Anchor>'
    '<x:PrintObject>False</x:PrintObject>'
    '<x:AutoFill>False</x:AutoFill>'
    '<x:FmlaMacro>[0]!{macro}</x:FmlaMacro>'
    '<x:TextHAlign>Center</x:TextHAlign>'
    '<x:TextVAlign>Center</x:TextVAlign>'
    '</x:ClientData>'
    '</v:shape>'
)
VML_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing"
VML_TYPE = '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'


def _vml(sheet_index, buttons):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    shapes = []
    for k, b in enumerate(buttons, start=1):
        letter, row = coordinate_from_string(b.cell)
        col = column_index_from_string(letter) - 1
        top = row - 1
        # Anchor: 左の列, 列内のずれ(px), 上の行, 行内のずれ, 右の列, ずれ, 下の行, ずれ（列と行は 0 始まり）
        anchor = f"{col}, 4, {top}, 3, {col + b.cols - 1}, 60, {top + b.rows - 1}, 17"
        shapes.append(VML_BUTTON.format(
            shape_id=sheet_index * 1024 + k, z=k, text=escape(b.text), macro=b.macro, anchor=anchor,
        ))
    return VML_HEAD.format(idmap=sheet_index) + "".join(shapes) + "</xml>"


def add_buttons(wb, path, buttons):
    """
    保存済みのブック（.xlsx / .xlsm）にボタンを足す。openpyxl はフォーム コントロールを書けないので、
    zip の中身を差し替える:
      - xl/drawings/vmlDrawingN.vml … ボタンの形と、押したときに動くマクロ（FmlaMacro）
      - シートの rels に vmlDrawing へのリレーション、シートの XML に <legacyDrawing>
    """
    if not buttons:
        return
    by_sheet = {}
    for b in buttons:
        by_sheet.setdefault(wb.sheetnames.index(b.sheet) + 1, []).append(b)
    files = {}
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        for item in infos:
            files[item.filename] = z.read(item.filename)
    for index, sheet_buttons in by_sheet.items():
        sheet_xml = f"xl/worksheets/sheet{index}.xml"
        rels_path = f"xl/worksheets/_rels/sheet{index}.xml.rels"
        vml_path = f"xl/drawings/vmlDrawing{index}.vml"
        files[vml_path] = _vml(index, sheet_buttons).encode()
        rel = f'<Relationship Id="rIdButtons" Type="{VML_REL}" Target="../drawings/vmlDrawing{index}.vml"/>'
        if rels_path in files:
            files[rels_path] = files[rels_path].decode().replace("</Relationships>", rel + "</Relationships>").encode()
        else:
            files[rels_path] = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + rel + "</Relationships>"
            ).encode()
        # legacyDrawing はシートの子要素の順番が決まっている（tableParts / extLst より前）
        legacy = '<legacyDrawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdButtons"/>'
        text = files[sheet_xml].decode()
        for tail in ("<tableParts", "<extLst", "</worksheet>"):
            if tail in text:
                text = text.replace(tail, legacy + tail, 1)
                break
        files[sheet_xml] = text.encode()
    types = files["[Content_Types].xml"].decode()
    if 'Extension="vml"' not in types:
        types = types.replace('<Default Extension="xml"', VML_TYPE + '<Default Extension="xml"', 1)
    files["[Content_Types].xml"] = types.encode()
    names = [i.filename for i in infos] + [n for n in files if n not in {i.filename for i in infos}]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in names:
            z.writestr(name, files[name])


def save_xlsm(wb, lecture_dir, name, buttons=()):
    """
    その節の完成コード（example/vba/*.bas）があれば、同じ中身のマクロ入りブック（.xlsm）も作る。
    openpyxl はマクロを書けないので、.xlsx を保存してから zip の中身を差し替える:
      - xl/vbaProject.bin を足す（tools/vba_project.py で組み立てたもの）
      - ブックの種類を「マクロ有効」に変え、vbaProject.bin へのリレーションを足す
      - ブックとシートにコード名（ThisWorkbook / Sheet1 …）を付ける（VBA の部品と対応させるため）
    """
    vba_dir = lecture_dir / "example" / "vba"
    bas_files = sorted(vba_dir.glob("*.bas"))
    if not bas_files:
        return
    wb.code_name = "ThisWorkbook"
    sheet_names = []
    for i, ws in enumerate(wb.worksheets, start=1):
        ws.sheet_properties.codeName = f"Sheet{i}"
        sheet_names.append(f"Sheet{i}")
    xlsm_path = lecture_dir / "book" / f"{name}.xlsm"
    tmp = lecture_dir / "book" / f"{name}.tmp.xlsx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tmp)
    vba = build_vba_project([(p.stem, p.read_text()) for p in bas_files], sheet_names)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(xlsm_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode()
                text = text.replace(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
                text = text.replace(
                    '<Default Extension="xml"',
                    '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default Extension="xml"',
                    1,
                )
                data = text.encode()
            elif item.filename == "xl/_rels/workbook.xml.rels":
                text = data.decode().replace(
                    "</Relationships>",
                    '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
                )
                data = text.encode()
            zout.writestr(item, data)
        zout.writestr("xl/vbaProject.bin", vba)
    tmp.unlink()
    add_buttons(wb, xlsm_path, buttons)
    print(f"[build_workbooks] {xlsm_path.relative_to(ROOT)}（マクロ入り）")


# 節ごとのボタン。表示する文字は本文（「◯◯」ボタンを押します）と揃える
SALES_BUTTONS = {
    "CalcAmounts": Button("売上", "L2", "金額を計算", "CalcAmounts"),
    "SumTotal": Button("売上", "L5", "売上合計を出す", "SumTotal"),
    "SumByCustomer": Button("売上", "L8", "顧客別に集計", "SumByCustomer"),
    "SumAllCustomers": Button("売上", "L11", "全顧客を集計", "SumAllCustomers"),
    "SumTotalWithBug": Button("売上", "L14", "合計（バグあり）", "SumTotalWithBug"),
}
DEMO_BUTTONS = {
    "ShowObjects": Button("売上", "L2", "ブックの情報を出す", "ShowObjects"),
    "DecorateSales": Button("売上", "L5", "表を整える", "DecorateSales"),
    "SplitByCustomer": Button("売上", "L8", "顧客別に分ける", "SplitByCustomer"),
    "DeleteDemoSheets": Button("売上", "L11", "増えたシートを消す", "DeleteDemoSheets"),
    "SaveBackup": Button("売上", "L14", "バックアップを保存", "SaveBackup"),
}
INVOICE_BUTTONS = {
    "CreateInvoices": Button("顧客", "F2", "請求書を作る", "CreateInvoices"),
    "DeleteInvoices": Button("顧客", "F5", "請求書シートを消す", "DeleteInvoices"),
}
FIRST_BUTTONS = {
    "Hello": Button("Sheet1", "D2", "あいさつする", "Hello"),
    "WriteCells": Button("Sheet1", "D5", "セルに書く", "WriteCells"),
    "CalcTaxIncluded": Button("Sheet1", "D8", "税込価格を計算", "CalcTaxIncluded"),
}


# ---------------------------------------------------------------------------
# 6 章（講義）: 表の形を見せる表計算デモ用のブック。マクロは入れず .xlsx だけを置く
# ---------------------------------------------------------------------------

def write_rows(ws, top, left, rows, widths=None):
    """1 行目を見出しにして rows を書く。None のセルは罫線だけ引く。"""
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = ws.cell(row=top + r, column=left + c, value=v)
            cell.border = BOX
            if r == 0:
                cell.font = BOLD
                cell.fill = HEADER_FILL
            elif isinstance(v, date):
                cell.number_format = DATE
            elif isinstance(v, int):
                cell.number_format = YEN
    for c, w in enumerate(widths or []):
        ws.column_dimensions[ws.cell(row=1, column=left + c).column_letter].width = w


def write_inputs(ws, refs):
    """デモで式を入れるセル（黄色）。"""
    for ref in refs:
        ws[ref].fill = INPUT_FILL
        ws[ref].border = BOX


def write_labels(ws, cells):
    for ref, text in cells.items():
        ws[ref] = text
        ws[ref].font = BOLD
        ws[ref].border = BOX


LONG_SALES = [
    # 月, 顧客, 商品, 数量（6-1 の「見やすい表」と同じ中身を 1 行に 1 件で並べたもの）
    ("10月", "ひまわり商事", "付箋", 5),
    ("10月", "ひまわり商事", "ノート", 2),
    ("10月", "ひまわり商事", "ホッチキス", 1),
    ("10月", "みなと物産", "付箋", 10),
    ("10月", "みなと物産", "ノート", 2),
    ("11月", "ひまわり商事", "付箋", 3),
    ("11月", "ひまわり商事", "ノート", 4),
    ("11月", "ひまわり商事", "ホッチキス", 2),
    ("11月", "みなと物産", "付箋", 6),
    ("11月", "みなと物産", "ノート", 1),
]


def build_table_design_61(path):
    wb = Workbook()
    # 見やすく整えた表。顧客名の結合・小計・空行・横に並ぶ月
    ws = wb.active
    ws.title = "見やすい表"
    write_rows(ws, 1, 1, [
        ["顧客", "商品", "10月", "11月"],
        ["ひまわり商事", "付箋", 5, 3],
        [None, "ノート", 2, 4],
        [None, "ホッチキス", 1, 2],
        ["小計", None, 8, 9],
        [None, None, None, None],
        ["みなと物産", "付箋", 10, 6],
        [None, "ノート", 2, 1],
        ["小計", None, 12, 7],
    ], widths=[16, 12, 8, 8])
    for rng in ("A2:A4", "A5:B5", "A7:A8", "A9:B9"):
        ws.merge_cells(rng)
    for row in (5, 9):
        for col in range(1, 5):
            ws.cell(row=row, column=col).fill = INPUT_FILL
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    for ref in ("A2", "A7"):
        ws[ref].alignment = Alignment(vertical="center")
    write_labels(ws, {"F1": "求めたいもの", "G1": "関数の答え", "H1": "正しい答え", "F2": "ひまわり商事の10月", "F3": "10月の合計"})
    ws["H2"], ws["H3"] = 8, 20
    write_inputs(ws, ["G2", "G3"])
    for ref in ("H2", "H3"):
        ws[ref].border = BOX
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 12

    # 1 行に 1 件に並べ直した表。同じ集計を関数で出す
    ws = wb.create_sheet("データ")
    write_rows(ws, 1, 1, [["月", "顧客", "商品", "数量"], *LONG_SALES], widths=[8, 16, 12, 8])
    write_labels(ws, {"F1": "求めたいもの", "G1": "関数の答え", "F2": "ひまわり商事の10月", "F3": "10月の合計"})
    write_inputs(ws, ["G2", "G3"])
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 12

    # データの表から、月を横に並べた「見せる表」を作る
    ws = wb.create_sheet("月別の表")
    write_rows(ws, 1, 1, [["月", "顧客", "商品", "数量"], *LONG_SALES], widths=[8, 16, 12, 8])
    write_labels(ws, {"F1": "顧客", "G1": "10月", "H1": "11月", "F2": "ひまわり商事", "F3": "みなと物産"})
    write_inputs(ws, ["G2", "H2", "G3", "H3"])
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 16
    save_plain(wb, path)


def build_masters_62(path):
    wb = Workbook()
    # 顧客名を手で打った売上。表記ゆれで SUMIF から漏れる
    ws = wb.active
    ws.title = "表記ゆれ"
    write_rows(ws, 1, 1, [
        ["日付", "顧客", "金額"],
        [date(2026, 10, 7), "ひまわり商事", 1360],
        [date(2026, 10, 8), "㈱ひまわり商事", 360],
        [date(2026, 10, 13), "ひまわり商事 ", 1100],
        [date(2026, 10, 15), "ひまわり商事", 6800],
        [date(2026, 10, 21), "ヒマワリ商事", 900],
    ], widths=[12, 18, 10])
    write_labels(ws, {"E1": "求めたいもの", "F1": "関数の答え", "E2": "ひまわり商事の合計", "E3": "全部の行の合計"})
    write_inputs(ws, ["F2", "F3"])
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 12

    # 売上には顧客コードだけ。顧客名と住所は、右の顧客の表から XLOOKUP で引く
    ws = wb.create_sheet("売上と顧客")
    write_rows(ws, 1, 1, [
        ["日付", "顧客コード", "顧客名", "住所"],
        [date(2026, 10, 2), "C002", None, None],
        [date(2026, 10, 6), "C003", None, None],
        [date(2026, 10, 7), "C001", None, None],
        [date(2026, 10, 7), "C001", None, None],
        [date(2026, 10, 8), "C003", None, None],
    ], widths=[12, 11, 18, 30])
    write_inputs(ws, [f"{col}{row}" for row in range(2, 7) for col in "CD"])
    write_rows(ws, 1, 6, [
        ["顧客コード", "顧客名", "住所"],
        ["C001", CUSTOMERS[0][0], CUSTOMERS[0][2]],
        ["C002", CUSTOMERS[1][0], CUSTOMERS[1][2]],
        ["C003", CUSTOMERS[2][0], CUSTOMERS[2][2]],
    ], widths=[11, 18, 30])
    ws.column_dimensions["E"].width = 3
    save_plain(wb, path)


def build_attendance_63(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "出欠"
    marks = {"E012": ["出席", "出席", "欠席"], "E027": ["出席", "遅刻", "出席"], "E031": ["出席", "出席", "出席"]}
    rows = [["日付", "社員番号", "研修", "出欠", "部署"]]
    for d, day in enumerate((1, 2, 3)):
        for emp, ms in marks.items():
            rows.append([date(2026, 4, day), emp, "新人研修", ms[d], None])
    write_rows(ws, 1, 1, rows, widths=[12, 10, 12, 8, 10])
    write_inputs(ws, [f"E{row}" for row in range(2, 11)])
    write_rows(ws, 1, 7, [
        ["社員番号", "氏名", "部署"],
        ["E012", "佐藤 花子", "営業部"],
        ["E027", "鈴木 一郎", "総務部"],
        ["E031", "田中 次郎", "営業部"],
    ], widths=[10, 12, 10])
    write_labels(ws, {"G6": "部署", "H6": "出席", "I6": "遅刻", "J6": "欠席", "G7": "営業部", "G8": "総務部"})
    write_inputs(ws, [f"{col}{row}" for row in (7, 8) for col in "HIJ"])
    ws.column_dimensions["F"].width = 3
    ws.column_dimensions["J"].width = 8
    save_plain(wb, path)


def save_plain(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"[build_workbooks] {path.relative_to(ROOT)}")


def pick(table, *names):
    return [table[n] for n in names]


def main():
    sales = make_sales()

    # 2 章: 空白のブックから始める。終えた状態として、空のシートにマクロを入れたブックを作る
    for lecture, buttons in (
        ("01-hello-macro", pick(FIRST_BUTTONS, "Hello")),
        ("02-cells-and-variables", pick(FIRST_BUTTONS, "Hello", "WriteCells", "CalcTaxIncluded")),
    ):
        wb = Workbook()
        wb.active.title = "Sheet1"
        save_xlsm(wb, SECTIONS / "02-first-macro" / lecture, "first-macro", buttons)

    # 3 章: 金額が空の売上明細 + 集計欄（受講者が関数と VBA で埋める）。.xlsx は 3-1〜3-3 で同じ
    # .xlsm は、その節のマクロを実行し終えた後の値（金額・合計・顧客別の合計）まで入れる
    totals = {}
    for _, customer, _, price, quantity in sales:
        totals[customer] = totals.get(customer, 0) + price * quantity

    def fill_amounts(ws):
        for r, (_, _, _, price, quantity) in enumerate(sales, start=2):
            ws.cell(row=r, column=6, value=price * quantity)

    def end_31(wb):
        fill_amounts(wb["売上"])

    def end_32(wb):
        ws = wb["売上"]
        fill_amounts(ws)
        ws["J2"] = sum(totals.values())
        ws["J5"] = totals[CUSTOMERS[0][0]]

    def end_34(wb):
        end_32(wb)
        for i, (name, *_rest) in enumerate(CUSTOMERS, start=5):
            wb["売上"].cell(row=i, column=10, value=totals[name])

    # 前の節のブックで続ける人もいるので、章で使うボタンは最初の節から全部置いておく
    summary = pick(SALES_BUTTONS, "CalcAmounts", "SumTotal", "SumByCustomer")
    for lecture, end_state, buttons, with_start in (
        ("01-calc-amounts", end_31, summary, True),
        ("02-sum-and-sumif", end_32, summary, False),
        ("03-debugger-intro", end_32, summary, False),
        ("04-debugging", end_34, summary + pick(SALES_BUTTONS, "SumAllCustomers", "SumTotalWithBug"), True),
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = "売上"
        write_sales_sheet(ws, sales, with_amount=False)
        write_summary_area(ws)
        save(wb, SECTIONS / "03-reproduce-functions" / lecture, "sales", end_state, buttons, with_start=with_start)

    # 4 章: 同じ表を金額まで埋めた状態（デモは集計済みの表から始める）
    for lecture, buttons in (
        ("01-object-model", pick(DEMO_BUTTONS, "ShowObjects")),
        ("02-sheet-operations", list(DEMO_BUTTONS.values())),
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = "売上"
        write_sales_sheet(ws, sales, with_amount=True)
        write_summary_area(ws)
        save(wb, SECTIONS / "04-what-vba-can-do" / lecture, "sales", buttons=buttons)

    # 5 章: 金額まで埋まった売上明細 + 顧客 + 請求書テンプレート
    # 「請求書シートを消す」は 5-3 で受講者が自分で置くボタンなので、5-3 を終えた .xlsm にだけ置く。
    # .xlsx は 5-1〜5-3 で同じ
    create = pick(INVOICE_BUTTONS, "CreateInvoices")
    for lecture, buttons, start_buttons, with_start in (
        ("01-copy-template", create, None, True),
        ("02-transfer-details", create, None, False),
        ("03-pdf-and-button", list(INVOICE_BUTTONS.values()), create, False),
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = "売上"
        write_sales_sheet(ws, sales, with_amount=True)
        write_customer_sheet(wb.create_sheet("顧客"))
        write_invoice_template(wb.create_sheet("請求書テンプレート"))
        save(wb, SECTIONS / "05-invoice-button" / lecture, "invoice", buttons=buttons, start_buttons=start_buttons, with_start=with_start)

    # 6 章: 講義の表計算デモ用。マクロは無いので .xlsx だけ
    chapter6 = SECTIONS / "06-table-design"
    build_table_design_61(chapter6 / "01-one-row-one-record" / "book" / "table-design.xlsx")
    build_masters_62(chapter6 / "02-master-tables" / "book" / "masters.xlsx")
    build_attendance_63(chapter6 / "03-design-new-sheet" / "book" / "attendance.xlsx")


if __name__ == "__main__":
    main()
